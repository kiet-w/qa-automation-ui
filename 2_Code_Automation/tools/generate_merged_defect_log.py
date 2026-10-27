import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
from PIL import Image as PILImage
import os

def build_merged_defect_log_with_images():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EduFAMS Defect Log"
    ws.views.sheetView[0].showGridLines = True

    font_family = "Segoe UI"
    
    # Header & Accent Fills
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    subtle_gray = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # Reporter Badges
    chan_fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
    chan_font = Font(name=font_family, size=10, bold=True, color="1967D2")
    
    kent_fill = PatternFill(start_color="FEF3D6", end_color="FEF3D6", fill_type="solid")
    kent_font = Font(name=font_family, size=10, bold=True, color="B06000")

    # Priority Fonts
    urgent_font = Font(name=font_family, size=10, bold=True, color="C5221F")
    high_font = Font(name=font_family, size=10, bold=True, color="D97706")
    medium_font = Font(name=font_family, size=10, color="2563EB")
    low_font = Font(name=font_family, size=10, color="4B5563")

    # General Fonts
    title_font = Font(name=font_family, size=15, bold=True, color="1F4E79")
    subtitle_font = Font(name=font_family, size=9.5, italic=True, color="4B5563")
    col_header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=9.5, color="111827")
    data_font_bold = Font(name=font_family, size=9.5, bold=True, color="111827")
    code_font = Font(name="Consolas", size=9, color="1E3A8A")

    # Borders
    thin_border_side = Side(style='thin', color='D1D5DB')
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Alignments
    center_align = Alignment(horizontal='center', vertical='center')
    center_top = Alignment(horizontal='center', vertical='top')
    left_top_wrap = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # -------------------------------------------------------------
    # Title Block (Rows 1 - 3)
    # -------------------------------------------------------------
    ws["A1"] = "EduFAMS Defect Log (Master QA Log with Embedded Evidence Screenshots)"
    ws["A1"].font = title_font
    ws.row_dimensions[1].height = 26

    ws["A2"] = "Consolidated Defect Tracking: Chan (DEF-001 to DEF-015) & Kent Tran (DEF-016 to DEF-023)  |  Screenshots directly embedded in column 'Evidence'"
    ws["A2"].font = subtitle_font
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 8

    # -------------------------------------------------------------
    # Headers (Row 4)
    # -------------------------------------------------------------
    headers = [
        "Defect ID",
        "Reported By",
        "Defect Summary",
        "Severity",
        "Priority",
        "Status",
        "Defect Type",
        "Environment",
        "Preconditions",
        "Steps to Reproduce",
        "Expected Result",
        "Actual Result",
        "Linked Test Case(s)",
        "Evidence (Screenshot)",
        "Notes"
    ]

    ws.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = col_header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # -------------------------------------------------------------
    # 23 Merged Defects (Chan: DEF-001 - DEF-015 | Kent Tran: DEF-016 - DEF-023)
    # -------------------------------------------------------------
    defects = [
        # --- CHAN'S DEFECTS (DEF-001 to DEF-015) ---
        (
            "DEF-001", "Chan", "Unable to click the Change button",
            "TBD", "High", "Open", "UI",
            "EduFAMS Admin Portal — School Users",
            "A school user row is displayed; screenshot shows USR-00002 (Bui Thi H).",
            "1. Open School Users. 2. Locate USR-00002 (Bui Thi H). 3. Click the pencil (Change/Edit) action.",
            "The selected school user's edit form opens.",
            "The pencil (Change/Edit) action cannot be clicked.",
            "TBD", "evidence/DEF-001_Evidence-01.png",
            "Screenshot highlights the pencil action for USR-00002."
        ),
        (
            "DEF-002", "Chan", "Unable to click the Delete button",
            "TBD", "High", "Open", "UI",
            "EduFAMS Admin Portal — School Users",
            "A school user row is displayed; screenshot shows USR-00002 (Bui Thi H), unassigned to a school.",
            "1. Open School Users. 2. Locate USR-00002 (Bui Thi H). 3. Click the trash (Delete) action.",
            "The delete confirmation or intended delete action is triggered.",
            "The Delete button cannot be clicked.",
            "TBD", "evidence/DEF-002_Evidence-01.png",
            "Screenshot highlights the Delete action for USR-00002."
        ),
        (
            "DEF-003", "Chan", "Import result column is labeled “Row” instead of “Number”",
            "TBD", "Low", "Open", "UI / Content",
            "EduFAMS Admin Portal — School Users — Bulk create",
            "The bulk import has completed and the Import result table is visible.",
            "1. Open School Users. 2. Go to Bulk create. 3. Import a CSV and review the Import result table.",
            "The column identifying each CSV entry is labeled “Number”.",
            "The column is labeled “Row”.",
            "TBD", "evidence/DEF-003_Evidence-01.png",
            "Screenshot shows the “Row” header and row indices in the import result table."
        ),
        (
            "DEF-004", "Chan", "Invalid Year value displays a generic validation message",
            "TBD", "Medium", "Open", "Validation / UX",
            "EduFAMS Admin Portal — Academic Years — Create",
            "The Create an academic year form is open.",
            "1. Open Academic Years and select Create academic year. 2. Enter an invalid value such as `wwe` in Year. 3. Enter valid start and end dates. 4. Select Create academic year.",
            "A clear, field-specific message explains the valid Year format and indicates the Year field needs correction.",
            "The form displays “Unable to create academic year” with the generic message System.ValidationFailed; it does not explain what is invalid in Year.",
            "TBD", "evidence/DEF-004_Evidence-01.png",
            "Screenshot shows `wwe` in Year and the generic validation error."
        ),
        (
            "DEF-005", "Chan", "Active class count is incorrect for FPT University — HCM City",
            "TBD", "High", "Open", "Data / Display",
            "EduFAMS Admin Portal — School management",
            "The school management list is displayed.",
            "1. Open School management. 2. Locate FPT University — HCM City (MOE code `FPT`). 3. Review its Active classes and Active students counts.",
            "Active classes displays the correct count for the school; confirm the expected count against the school's class data.",
            "Active classes displays 0 while Active students displays 30.",
            "TBD", "evidence/DEF-005_Evidence-01.png",
            "Screenshot highlights the Active classes value 0 for FPT. Expected count not provided."
        ),
        (
            "DEF-006", "Chan", "Search students field is missing placeholder text",
            "TBD", "Low", "Open", "UI / Content",
            "EduFAMS Admin Portal — Students",
            "The Students page is open.",
            "1. Open Students. 2. Inspect the Search students input while it is empty.",
            "The search field displays placeholder text that helps users understand what they can search for.",
            "The empty Search students field has no visible placeholder text.",
            "TBD", "evidence/DEF-006_Evidence-01.png",
            "Screenshot shows the blank search input; expected placeholder wording not specified."
        ),
        (
            "DEF-007", "Chan", "Fee batch search does not filter results",
            "TBD", "Medium", "Open", "Functional",
            "EduFAMS Admin Portal — Fee Batches",
            "Fee publication batches are listed.",
            "1. Open Fee Batches. 2. Enter `GOPs` in Search fee batches. 3. Select Search.",
            "Results are filtered to batches matching `GOPs`.",
            "Unrelated batch titles remain visible alongside the matching `GOPs` result.",
            "TBD", "evidence/DEF-007_Evidence-01.png",
            "Screenshot shows `GOPs` in the search field while `ffff332`, `Học Phí 2026`, and `ffff` remain in the results."
        ),
        (
            "DEF-008", "Chan", "Unable to create fee publication when no active academic year exists",
            "TBD", "Medium", "Open", "Functional / Configuration",
            "EduFAMS Admin Portal — Fee Batches — Create fee publication",
            "The Create fee publication form is open; no active academic year is configured.",
            "1. Open Fee Batches and start creating a fee publication. 2. Enter the publication details and targets. 3. Submit the form.",
            "The publication is created when valid required details are provided and an active academic year is configured, or the prerequisite is clearly explained before submission.",
            "Creation fails with “Could not create fee publication” and “No active academic year.”",
            "TBD", "evidence/DEF-008_Evidence-01.png",
            "Screenshot shows the error with title `SOP900` and period/charge/due dates entered."
        ),
        (
            "DEF-009", "Chan", "Fee publication batches page fails to load data",
            "TBD", "Urgent", "Open", "Backend / Data",
            "EduFAMS Admin Portal — Fee Batches",
            "The Fee publication batches page is opened.",
            "1. Open Fee Batches. 2. Wait for the publication batches list to load.",
            "Fee batch data loads and the user can review the list.",
            "The page shows “Unable to load data” with error `Invalid column name 'school_id'`.",
            "TBD", "evidence/DEF-009_Evidence-01.png",
            "The page offers a “Try again” action, but the displayed database column error prevents data from loading."
        ),
        (
            "DEF-010", "Chan", "Fee management page fails to load fee types",
            "TBD", "Urgent", "Open", "Functional / Data",
            "EduFAMS Admin Portal — Fee management",
            "The Fee management page is opened.",
            "1. Open Fees. 2. Wait for fee type data to load.",
            "Fee types and related fee management data are displayed.",
            "The page shows “Unable to load data” and “Fee information could not be loaded. Try again.”",
            "TBD", "evidence/DEF-010_Evidence-01.png",
            "Screenshot shows the fee types search area and the load error with a “Try again” action."
        ),
        (
            "DEF-011", "Chan", "Student import template shows the wrong date format",
            "TBD", "Medium", "Open", "UI / Content",
            "EduFAMS Admin Portal — Students — Bulk import template",
            "The student bulk import template is open in a spreadsheet application.",
            "1. Download and open the student bulk import template. 2. Review the Date of Birth and Enrollment Start column headers.",
            "Both date columns specify the expected format mm/dd/yyyy.",
            "The headers specify dd/mm/yyyy.",
            "TBD", "evidence/DEF-011_Evidence-01.png",
            "Screenshot shows dd/mm/yyyy for Date of Birth and Enrollment Start; expected format supplied as mm/dd/yyyy."
        ),
        (
            "DEF-012", "Chan", "Unable to import students using the bulk student upload",
            "TBD", "Urgent", "Open", "Functional / Validation",
            "EduFAMS Admin Portal — Students — Bulk creation",
            "The Bulk student creation form is open and a student CSV file is selected.",
            "1. Open Students and select Bulk. 2. Choose the completed student import CSV. 3. Review the validation result.",
            "A CSV using the provided student import template is accepted and proceeds to row review.",
            "The upload is rejected with “Invalid student import template” and a message listing expected columns: Display Name, Student Nric, Date of Birth, Student Citizenship, Enrollment Start, and Class Code.",
            "TBD", "evidence/DEF-012_Evidence-01.png",
            "Screenshot shows Bulk_Students_Import_Template (1).csv selected and Submit students disabled. Priority marked Urgent per user."
        ),
        (
            "DEF-013", "Chan", "Bulk class preview column is labeled “Row” instead of “No”",
            "TBD", "Low", "Open", "UI / Content",
            "EduFAMS Admin Portal — Classes — Bulk creation",
            "A class import CSV has been uploaded and the Preview and edit table is displayed.",
            "1. Open Classes and select Bulk. 2. Upload a class template CSV. 3. Review the Preview and edit table.",
            "The row-number column is labeled “No”.",
            "The column is labeled “Row”.",
            "TBD", "evidence/DEF-013_Evidence-01.png",
            "Screenshot shows the “Row” header and row numbers 2 and 3. Expected label supplied as “No”."
        ),
        (
            "DEF-014", "Chan", "Academic level in bulk class preview is a text field instead of a dropdown",
            "TBD", "Medium", "Open", "UI / Usability",
            "EduFAMS Admin Portal — Classes — Bulk creation",
            "A class import CSV has been uploaded and the Preview and edit table is displayed.",
            "1. Open Classes and select Bulk. 2. Upload a class template CSV. 3. Review the Academic level control for a row.",
            "Academic level is selected from a dropdown containing available academic levels.",
            "Academic level is presented as a free-text input.",
            "TBD", "evidence/DEF-014_Evidence-01.png",
            "Screenshot shows text inputs containing Primary 1 and Primary 2 in the Academic level column."
        ),
        (
            "DEF-015", "Chan", "Student profile does not display NRIC information",
            "TBD", "High", "Open", "Data / Display",
            "EduFAMS Admin Portal — Students — Student profile",
            "A student profile is open.",
            "1. Open Students. 2. Open the student profile (screenshot shows OP1). 3. Review the displayed student details.",
            "The student's NRIC information is displayed in the profile, subject to the user's access permissions.",
            "No NRIC field or value is visible in the profile details.",
            "TBD", "evidence/DEF-015_Evidence-01.png",
            "Screenshot shows student name, date of birth, category, school, class, and enrollment details, with no NRIC information."
        ),

        # --- KENT TRAN'S DEFECTS (DEF-016 to DEF-023) ---
        (
            "DEF-016", "Kent Tran", "MOEFAS page fails to load schemes with HTTP 404",
            "Blocker", "Urgent", "Open", "Backend / API",
            "EduFAMS Admin Portal — FINANCIAL — MOEFAS",
            "Logged in as MOE HQ Administrator (admin@IAM.edu.vn).",
            "1. Open Admin Portal. 2. Under FINANCIAL in sidebar, select MOEFAS (/moefas). 3. Review page load and network response.",
            "The page loads available MOEFAS schemes from backend or renders a clean empty state with a 'New scheme' button (HTTP 200).",
            "API call GET /api/fas/v1/fas/moe-schemes fails with HTTP 404. Page displays red error alert: “Unable to load data: The MOEFAS schemes could not be loaded. Please try again.”",
            "TS-R4-FAS-001", "evidence/DEF-016_Evidence-01.png",
            "Backend controller route unmapped or Vite proxy misconfigured for /api/fas/. Fatal client demo breaker."
        ),
        (
            "DEF-017", "Kent Tran", "Navigation sidebar exposes unbuilt scaffold placeholder pages and dashboard header typo",
            "Major", "High", "Open", "UI / Content",
            "EduFAMS Admin Portal — Top-up, FAS & Root Dashboard",
            "Admin portal is opened under any administrative role.",
            "1. Log in to Admin Portal and inspect root header. 2. Select EDUSAVE -> Top-up. 3. Select FINANCIAL -> FAS.",
            "Dashboard displays production widgets with correct spelling 'Financial'. Unreleased modules are hidden from sidebar navigation.",
            "Root dashboard header displays “Education Fanancial Account Management System” (spelling error) and body says “This feature is scaffolded...”. Both /top-up and /fas display blank scaffold cards.",
            "TS-R3-DASH-001", "evidence/DEF-017_Evidence-01.png",
            "Needs typo fix in translation.json ('Fanancial' -> 'Financial') and feature-flag hiding in manifests."
        ),
        (
            "DEF-018", "Kent Tran", "MOE HQ role is bounced with Access Restricted error upon accessing Class Management",
            "Major", "Medium", "Open", "Routing / Security Guard",
            "EduFAMS Admin Portal — Class Management",
            "Logged in as MOE HQ Administrator (admin@IAM.edu.vn).",
            "1. Log in as MOE HQ Admin. 2. Navigate directly to /class-management (or click class link from school overview).",
            "MOE HQ is granted read-only auditing access across all school classes, or presented with an informative 403 / role switch screen.",
            "Router guard blocks access, flashes red toast “Access restricted: Your account does not have permission...”, and bounces user to root dashboard.",
            "TS-R3-F04-001", "evidence/DEF-018_Evidence-01.png",
            "class-management/manifest.ts restricts access to SCHOOL only, trapping MOE HQ users."
        ),
        (
            "DEF-019", "Kent Tran", "Fee Batches table and detail view render raw technical GUIDs instead of human-readable names",
            "Moderate", "Medium", "Open", "UI / Data Display",
            "EduFAMS Admin Portal — Fee Batches",
            "Completed or draft fee publication batches exist in the system.",
            "1. Open Fee Batches. 2. Review Fee Component column. 3. Select View details on a batch. 4. Review Fee Component and Created by fields.",
            "Fee Component displays friendly component name (e.g., “Standard Miscellaneous Fee”) and Created by displays admin name/email.",
            "Table and details display raw GUIDs ffff0001-0000-0000-0000-000000000002 and 01A159A5-0F27-.... Details view also lacks student charges sub-table.",
            "TS-R3-F12-003", "evidence/DEF-019_Evidence-01.png",
            "Requires entity lookup/join against fee component and user tables before rendering."
        ),
        (
            "DEF-020", "Kent Tran", "Newly created student displays raw database UUID instead of standard STU-xxx format",
            "Moderate", "Low", "Open", "Data / Display",
            "EduFAMS Admin Portal — Students",
            "The Create student form is completed and submitted.",
            "1. Open Students and select Create student. 2. Create student Lucas Tan Wei Jie. 3. Review the student management table and student profile header.",
            "System assigns a standardized business code (e.g. STU-00005) consistent with baseline records (STU-001, STU-003).",
            "Table and profile header display full database UUID Lucas Tan Wei Jie f73e9aa4-0143-4826-9e48-93e184200d07.",
            "TS-R3-F05-002", "evidence/DEF-020_Evidence-01.png",
            "Needs sequential business code generator in StudentSchoolService on record persistence."
        ),
        (
            "DEF-021", "Kent Tran", "Database contains invalid test mock student (Age 25 in Primary 2 with corrupted class code)",
            "Moderate", "Low", "Open", "Data / Sanitation",
            "EduFAMS Admin Portal — Students",
            "Logged in as School User for Rosyth Primary School (staff1@IAM.edu.vn).",
            "1. Open Students. 2. Inspect student row for Trần Tuấn Kiệt. 3. Review Date of Birth, Category, and Class.",
            "All mock student records reflect valid primary school age limits (Ages 6–12) and valid class names.",
            "Record shows DOB May 18, 2001 (Age 25 in primary school) assigned to class 111111111111111111 · P2-X.",
            "TS-R3-F05-DATA", "evidence/DEF-021_Evidence-01.png",
            "Dirty test mock data in StudentSchoolDatabase.dbo.students. Needs purge script."
        ),
        (
            "DEF-022", "Kent Tran", "Parent portal child profile card displays truncated age label missing numerical value",
            "Minor", "Low", "Open", "UI / Content",
            "EduFAMS Parent Portal E-Service — Dashboard",
            "Logged in as verified parent via SingPass MockPass (e.g., S7890123F - NGUYEN VAN MINH).",
            "1. Access Parent Portal (http://localhost:5174/). 2. Authenticate as parent. 3. Review the child Education Account card header.",
            "Card header displays formatted age integer (e.g., EDU-00172 · Age 11 · Active).",
            "Numerical age is omitted, literally displaying EDU-00172 · Age Active.",
            "TS-R9-PARENT-001", "evidence/DEF-022_Evidence-01.png",
            "Missing age integer in React JSX template string interpolation."
        ),
        (
            "DEF-023", "Kent Tran", "FAS application submitted by parent has no review or approval queue in Admin Portal",
            "Major", "High", "Open", "Functional / Workflow",
            "EduFAMS Parent Portal & Admin Portal",
            "Parent has submitted a financial assistance application on Parent Portal (/fas/apply).",
            "1. Complete and submit FAS application on Parent Portal. 2. Log in to Admin Portal as School Clerk or MOE HQ. 3. Attempt to locate incoming application for verification.",
            "Submitted applications appear in an administrative queue for review, document verification, and approval/rejection.",
            "Admin Portal /fas is a blank placeholder card and /moefas fails with 404; submitted applications cannot be audited or actioned via UI.",
            "TS-R4-FAS-FLOW", "evidence/DEF-023_Evidence-01.png",
            "Broken end-to-end integration between apps/parent-portal intake and apps/admin-portal back-office review."
        )
    ]

    # Target thumbnail image dimensions inside Column N (Evidence)
    # Column N width = 38 (~270px)
    # Row height = 110 points (~146px)
    # Target image size: max width 250px, max height 135px
    for row_idx, row_data in enumerate(defects, start=5):
        ws.row_dimensions[row_idx].height = 110
        bg_fill = zebra_fill if row_idx % 2 == 0 else white_fill
        
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = bg_fill
            cell.border = thin_border
            cell.alignment = left_top_wrap
            
            # Formatting specifics
            if col_idx == 1:  # Defect ID
                cell.font = code_font
                cell.alignment = center_top
            elif col_idx == 2:  # Reported By
                cell.alignment = center_top
                if val == "Chan":
                    cell.fill = chan_fill
                    cell.font = chan_font
                else:
                    cell.fill = kent_fill
                    cell.font = kent_font
            elif col_idx == 4:  # Severity
                cell.alignment = center_top
            elif col_idx == 5:  # Priority
                cell.alignment = center_top
                if val == "Urgent": cell.font = urgent_font
                elif val == "High": cell.font = high_font
                elif val == "Medium": cell.font = medium_font
                elif val == "Low": cell.font = low_font
            elif col_idx == 6:  # Status
                cell.alignment = center_top
                cell.font = data_font_bold
            elif col_idx == 7:  # Defect Type
                cell.alignment = center_top
            elif col_idx == 13: # Linked Test Case
                cell.font = code_font
                cell.alignment = center_top
            elif col_idx == 14: # Evidence Column
                # Clear text so image overlays cleanly, or keep relative path in background
                cell.value = ""
                img_path = val  # e.g. "evidence/DEF-001_Evidence-01.png"
                if os.path.exists(img_path):
                    try:
                        with PILImage.open(img_path) as pil_im:
                            orig_w, orig_h = pil_im.size
                        ratio = min(250.0 / orig_w, 135.0 / orig_h)
                        target_w = int(orig_w * ratio)
                        target_h = int(orig_h * ratio)
                        
                        excel_img = OpenpyxlImage(img_path)
                        excel_img.width = target_w
                        excel_img.height = target_h
                        
                        # Add image directly anchored to cell N{row_idx}
                        anchor_pos = f"N{row_idx}"
                        ws.add_image(excel_img, anchor_pos)
                    except Exception as e:
                        print(f"Error embedding {img_path}: {e}")
                        cell.value = val
                else:
                    cell.value = val
                    cell.font = code_font

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:O{4 + len(defects)}"

    col_widths = {
        1: 12,  # Defect ID
        2: 14,  # Reported By
        3: 36,  # Defect Summary
        4: 12,  # Severity
        5: 12,  # Priority
        6: 10,  # Status
        7: 18,  # Defect Type
        8: 26,  # Environment
        9: 30,  # Preconditions
        10: 40, # Steps to Reproduce
        11: 34, # Expected Result
        12: 36, # Actual Result
        13: 16, # Linked Test Case(s)
        14: 38, # Evidence (Screenshot Image Embedded!)
        15: 35  # Notes
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output_path = "docs/EduFAMS_Merged_Defect_Log.xlsx"
    wb.save(output_path)
    print(f"Merged Defect Log with directly embedded images generated successfully at {output_path}")

if __name__ == "__main__":
    build_merged_defect_log_with_images()
