"""Publish one actual UAT-007 run into the R08 UAT evidence area."""
import copy
import html
import json
import os
from pathlib import Path
import re
import shutil
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports/fs08/fs08-uat007-contact-update"
TARGET = ROOT / "reports/uat/r08/fs08-uat007-contact-update"
MASTER = ROOT / "reports/uat/r08/TEST_EXECUTION_EVIDENCE_R08_UAT_MASTER.xlsx"
CASES = ROOT / "data/fs08/uat007_coverage_cases.json"
SPEC = ROOT / "testcases/fs08/fs08-uat007-contact-update/TC_FS08_UAT007_Contact_Update.md"
HEADERS = ["ID", "Scenario ID", "Title", "Test Suite", "Page", "Severity", "Priority",
           "Execution Status", "Test Type", "Environment", "Preconditions", "Test Data",
           "Test Steps", "Expected Result", "Actual Result", "Assumption Ref", "Defect ID",
           "Traceability", "Evidence"]
def original_cases():
    rows = {}
    for line in SPEC.read_text(encoding="utf-8").splitlines():
        if line.startswith("| `TC-UAT007-"):
            cells = [part.strip().replace("`", "").replace("<br>", "\n") for part in line.strip().split("|")[1:-1]]
            if len(cells) == 19:
                rows[cells[0]] = dict(zip(HEADERS, cells))
    return rows
def actual(test):
    steps = test.get("metadata", {}).get("steps", [])
    if steps:
        result = next((step.get("actual") for step in reversed(steps) if step.get("actual")), "")
        if result:
            return result[:500]
    errors = re.findall(r"^E\s+(.+)$", test.get("call", {}).get("longrepr", ""), re.M)
    return (errors[-1] if errors else test.get("outcome", "Unknown"))[:500]
def status(test):
    if test is None:
        return "Not Run"
    if test.get("outcome") == "passed":
        return "Passed"
    if test.get("outcome") == "skipped":
        return "Skipped"
    if not test.get("metadata", {}).get("steps") and "TimeoutError" in test.get("call", {}).get("longrepr", ""):
        return "Blocked"
    return "Failed"
def screenshot(tc_id):
    token = tc_id.lower().replace("-", "_")
    matches = list((SOURCE / "screenshots").glob(f"*{token}*/actual_result.png"))
    return max(matches, key=lambda path: path.stat().st_mtime) if matches else None
def results():
    report = json.loads((SOURCE / "report.json").read_text(encoding="utf-8"))
    found = {}
    for test in report.get("tests", []):
        match = re.search(r"TC-UAT007-\d+", test.get("nodeid", ""), re.I)
        if match:
            found[match.group().upper()] = test
    rows = original_cases()
    for case in json.loads(CASES.read_text(encoding="utf-8")):
        inp = case["input"]
        rows[case["id"]] = {
            "ID": case["id"], "Scenario ID": "TS-UAT007-02", "Title": case["title"],
            "Test Suite": f'Atomic {case["field"]}', "Page": "/profile Edit", "Severity": "S1",
            "Priority": "High", "Test Type": "Atomic validation", "Environment": "localhost:5174 + MockPass",
            "Preconditions": "Howard authenticated; profile edit mode", "Test Data": json.dumps(inp, ensure_ascii=False),
            "Test Steps": f'Open /profile; edit {case["field"]}; enter fixture; blur; inspect DOM state',
            "Expected Result": f'{case["kind"]}: {case.get("message", "documented field boundary/normalization")}',
            "Assumption Ref": "", "Defect ID": "", "Traceability": case["source"], "Evidence": "",
        }
    ordered = [rows[key] for key in sorted(rows, key=lambda value: int(value.rsplit("-", 1)[1]))]
    for row in ordered:
        test = found.get(row["ID"])
        row["Execution Status"] = status(test)
        row["Actual Result"] = actual(test) if test else "Not executed in this run"
        row["Evidence"] = screenshot(row["ID"])
    return report.get("run_id", "RUN-UNKNOWN"), ordered
def write_ticket_book(run_id, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "UAT-007"
    ws.append(["UAT-007 Parent Contact Update"])
    counts = Counter(row["Execution Status"] for row in rows)
    ws.append([f'{run_id} | {len(rows)} TCs | {counts["Passed"]} Passed | {counts["Failed"]} Failed | {counts["Blocked"]} Blocked'])
    ws.append([])
    ws.append(HEADERS)
    for row in rows:
        ws.append([str(row.get(key, "")) if key != "Evidence" else "" for key in HEADERS])
        index = ws.max_row
        shot = row["Evidence"]
        if shot:
            link = ws.cell(index, 19, "📷 Actual Result")
            link.hyperlink = os.path.relpath(shot, TARGET)
            link.style = "Hyperlink"
        ws.cell(index, 8).fill = PatternFill("solid", fgColor={"Passed":"DCFCE7", "Failed":"FEE2E2", "Blocked":"FEF3C7"}.get(row["Execution Status"], "F1F5F9"))
        for cell in ws[index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:S{ws.max_row}"
    for col, width in {"A":18,"B":18,"C":42,"D":24,"E":19,"H":18,"L":30,"M":48,"N":48,"O":55,"R":36,"S":21}.items():
        ws.column_dimensions[col].width = width
    TARGET.mkdir(parents=True, exist_ok=True)
    out = TARGET / "TEST_EXECUTION_EVIDENCE_FS08_UAT007.xlsx"
    wb.save(out)
    return counts
def write_master(run_id, rows, counts):
    wb = load_workbook(MASTER)
    ws = wb.active
    positions = [i for i in range(5, ws.max_row + 1) if ws.cell(i, 2).value == "UAT-007"]
    start = min(positions)
    template = [copy.copy(ws.cell(start + 1, col)._style) for col in range(1, 20)]
    for i in reversed(positions):
        ws.delete_rows(i)
    ws.insert_rows(start, len(rows))
    for offset, row in enumerate(rows):
        values = [None,"UAT-007",row["ID"],row.get("Assumption Ref",""),row["Scenario ID"],row["Title"],
                  row["Test Suite"],row["Page"],row["Severity"],row["Priority"],row["Execution Status"],
                  row["Test Type"],row["Environment"],row["Preconditions"],row["Test Data"],row["Test Steps"],
                  row["Expected Result"],row["Actual Result"],row["Traceability"]]
        for col, value in enumerate(values, 1):
            cell = ws.cell(start + offset, col, value)
            cell._style = copy.copy(template[col - 1])
        shot = row["Evidence"]
        if shot:
            ws.cell(start + offset, 18).hyperlink = os.path.relpath(shot, MASTER.parent)
    for i in range(5, ws.max_row + 1):
        ws.cell(i, 1, i - 4)
    ws["A2"] = f'Latest UAT-007 {run_id}: {counts["Passed"]}/{len(rows)} Passed, {counts["Failed"]} Failed, {counts["Blocked"]} Blocked. Other modules historical; not rerun.'
    wb.save(MASTER)
def main():
    run_id, rows = results()
    counts = write_ticket_book(run_id, rows)
    write_master(run_id, rows, counts)
    shutil.copy2(SOURCE / "report.json", TARGET / "report.json")
    detail = "\n".join(f'<tr><td>{html.escape(r["ID"])}</td><td>{html.escape(r["Title"])}</td><td>{html.escape(r["Execution Status"])}</td><td>{html.escape(r["Actual Result"])}</td></tr>' for r in rows)
    (TARGET / "execution_report.html").write_text(f'<!doctype html><meta charset="utf-8"><title>UAT-007 {run_id}</title><style>body{{font:14px Arial;max-width:1400px;margin:30px auto}}td,th{{border:1px solid #ddd;padding:8px;vertical-align:top}}table{{border-collapse:collapse;width:100%}}</style><h1>UAT-007 — {run_id}</h1><p>{counts["Passed"]} Passed · {counts["Failed"]} Failed · {counts["Blocked"]} Blocked / {len(rows)} TCs</p><p><a href="../../../fs08/fs08-uat007-contact-update/execution_report.html">Technical report and screenshots</a> · <a href="TEST_EXECUTION_EVIDENCE_FS08_UAT007.xlsx">Single-tab evidence</a></p><table><tr><th>TC</th><th>Scenario</th><th>Status</th><th>Actual result</th></tr>{detail}</table>', encoding="utf-8")
    print(run_id, dict(counts), TARGET)
if __name__ == "__main__":
    main()
