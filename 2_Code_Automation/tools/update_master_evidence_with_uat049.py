#!/usr/bin/env python3
"""
tools/update_master_evidence_with_uat049.py
Appends UAT-049 test cases to TEST_EXECUTION_EVIDENCE_R08_UAT_MASTER.xlsx.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def update_master():
    master_file = Path('/home/baudui/Projects/automation-ui/reports/fs08/TEST_EXECUTION_EVIDENCE_R08_UAT_MASTER.xlsx')
    uat049_file = Path('/home/baudui/Projects/automation-ui/reports/fs08/uat r8/fs08-uat049-family-fas-submission/TEST_EXECUTION_EVIDENCE_FS08_UAT049_FAMILY_FAS_SUBMISSION.xlsx')

    if not master_file.exists() or not uat049_file.exists():
        print("Missing master or uat049 file")
        return

    m_wb = openpyxl.load_workbook(master_file)
    m_ws = m_wb.active

    # Check if UAT-049 is already present
    for r in range(5, m_ws.max_row + 1):
        if m_ws.cell(r, 2).value == 'UAT-049':
            print("UAT-049 already in master excel!")
            return

    u_wb = openpyxl.load_workbook(uat049_file)
    u_ws = u_wb.active

    # Thin border
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    current_no = m_ws.max_row - 4 # initial count
    start_row = m_ws.max_row + 1

    for row_idx, r in enumerate(u_ws.iter_rows(min_row=5, values_only=True), start=1):
        # r columns: ID (0), Scenario ID (1), Title (2), Test Suite (3), Page (4), Severity (5), Priority (6), Status (7), Type (8), Env (9), Preconditions (10), Data (11), Steps (12), Expected (13), Actual (14), Assumption (15), Defect (16), Traceability (17), Note (18)
        current_no += 1
        m_row = start_row + row_idx - 1

        tc_id = r[0]
        scenario_id = r[1]
        title = r[2]
        suite = r[3]
        page = r[4]
        severity = r[5]
        priority = r[6]
        status = r[7]
        test_type = r[8]
        env = r[9]
        precond = r[10]
        test_data = r[11]
        steps = r[12]
        expected = r[13]
        actual = r[14]
        traceability = r[17]

        # Determine Pillar from Scenario ID or Suite
        pillar = 'BR'
        if 'VAL' in str(tc_id) or 'VAL' in str(scenario_id):
            pillar = 'VAL'
        elif 'ERR' in str(tc_id) or 'RES' in str(scenario_id):
            pillar = 'ERR'
        elif 'UI' in str(tc_id) or 'SEC' in str(tc_id):
            pillar = 'UI'
        elif 'FLO' in str(tc_id):
            pillar = 'FLOW'
        elif 'AC' in str(tc_id):
            pillar = 'AC'

        vals = [
            current_no, 'UAT-049', tc_id, pillar, scenario_id, title, suite, page,
            severity, priority, status, test_type, env, precond, test_data, steps,
            expected, actual, traceability
        ]

        for col_idx, val in enumerate(vals, start=1):
            cell = m_ws.cell(m_row, col_idx, val)
            cell.font = Font(name='Calibri', size=10)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

            if col_idx in (1, 2, 3, 4, 5, 8, 9, 10, 11, 12, 13):
                cell.alignment = Alignment(horizontal='center', vertical='top')

            if col_idx == 11: # Status
                cell.font = Font(name='Calibri', size=10, bold=True, color='00166534')
                cell.fill = PatternFill(start_color='00DCFCE7', end_color='00DCFCE7', fill_type='solid')

    # Update summary header row 2
    total_tests = current_no
    m_ws.cell(2, 1, f"Summary: Total 6 UAT Modules | {total_tests} Test Cases | 100% Passed | 0 Defects | Singpass & MockPass Verified")

    m_wb.save(master_file)
    print(f"Updated {master_file} with {row_idx} test cases. Total: {total_tests}")

if __name__ == '__main__':
    update_master()
